# Create your views here.
import datetime
import os
from ast import literal_eval
from decimal import Decimal
from pprint import pprint
from sqlite3 import Date
from datetime import date

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from io import BytesIO
import base64
import pandas as pd
from django.contrib import admin
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import PasswordChangeForm
from django.contrib.auth.mixins import PermissionRequiredMixin
from django.contrib.auth.models import Permission
from django.contrib.sessions.models import Session
from django.core import serializers
from django.core.cache import cache
from django.core.files.storage import FileSystemStorage
from django.core.paginator import Paginator
from django.db.models import Q
from django.forms import model_to_dict
from django.http import HttpResponse, JsonResponse, HttpResponseRedirect
from django.shortcuts import redirect, render, get_object_or_404
from django.template.backends.django import Template
from django.urls import reverse
from django.utils import timezone
from django.utils.decorators import method_decorator
from django.utils.text import slugify
from django.utils.translation import gettext as _
from django.views.generic import TemplateView
from django_dump_die.middleware import dd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from datetime import datetime, timezone
from django.db.models import Sum, Q, ExpressionWrapper, F, DurationField, Max
from django.utils.timezone import now
from datetime import timedelta
from django.db.models import Sum
import tempfile
import os
from django.core.files import File

from configurations.helper_config import verify_sql_query
from configurations.models import ActionLog, Prescripteur, PrescripteurPrestataire, Prestataire, Specialite, Secteur, \
    Bureau,TypeActe,BusinessUnit,Branche,Banque,Affection,Apporteur,ApporteurInternational,CategorieAffection,Devise,\
    TypePrestataire, User, AuthGroup, TypeEtablissement,Tarif, Rubrique, RegroupementActe, Acte, ReseauSoin, \
    PrestataireReseauSoin, WsBoby, ParamWsBoby, Affection, BackgroundQueryTask, ParamProduitCompagnie, Compagnie, \
    AlimentMatricule, ParamActe, TypeApporteur, TypePersonne, Pays, TypeCompagnie, TypeGarant, RisqueProduit, Carosserie, \
    CategorieVehicule, Civilite, CompteTresorerie, ConditionsAssurance, Carburant, Formule, Fractionnement, Garantie, GarantieFormule, \
    Groupe, ModeReglement
from inov import settings
# Create your views here.
from production.models import TarifPrestataireClient, Client, Aliment, AlimentFormule, Mouvement, MouvementAliment, \
    Carte, Quittance, Reglement, Courrier, Produit, PoliceAssureur, Police, HistoriquePolice, MouvementPolice
from analysecontrole.models import AnalysePortefeuille
from production.templatetags.my_filters import money_field, convertir_date_multiformat
from shared.enum import PasswordType, Statut, StatutValidite, BaseCalculTM, StatutPaiementSinistre, TypePortefeuille, \
    SatutBordereauDossierSinistres, StatutSinistre

from production.templatetags.my_filters import money_field, convertir_date_multiformat, supprimer_espaces, convertir_date_jj_mm_aaaa, format_montant, money_format_mille, \
    rendre_html


class AnalysePortefeuilleView(PermissionRequiredMixin,TemplateView):
    template_name = 'analyse/analyse.html'
    permission_required = "analysecontrole.view_analyseportefeuille"
    model = AnalysePortefeuille

    def get(self, request, *args, **kwargs):
        context_original = self.get_context_data(**kwargs)

        analyseportefeuille = AnalysePortefeuille.objects.all().order_by('-id')

        today = datetime.now(tz=timezone.utc)
        compagnies = Compagnie.objects.order_by('nom')
        businessunit = BusinessUnit.objects.order_by('libelle')

        commercials = []
        utilisateur = User.objects.all().order_by('-first_name').exclude(is_admin_group=1)
        for user in utilisateur:
            if user.is_commercial:
                commercials.append(user)

        context_perso = {'analyseportefeuille': analyseportefeuille, 'compagnies': compagnies, 'businessunit': businessunit, 'commercials': commercials, 'today': today}

        context = {**context_original, **context_perso}

        return self.render_to_response(context)

    def post(self):
        pass

    def get_context_data(self, **kwargs):
        pprint(kwargs)
        return {
            **super().get_context_data(**kwargs),
            **admin.site.each_context(self.request),
            "opts": self.model._meta,
        }


# Portefeuille par compagnie
def generate_excel_portefeuille_compagnie(compagnies, date_requete):
    """Génère un fichier Excel unique regroupant les portefeuilles de toutes les compagnies."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portefeuille"

    # En-tête du fichier
    sheet.append(["", "DATE DE LA REQUÊTE", date_requete])
    sheet.append([])  # Ligne vide

    headers = [
        "POLICE", "CLIENT", "TYPE DE CLIENT", "BRANCHE", "PRODUIT", "ÉCHÉANCE",
        "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N", "STATUT"
    ]

    for compagnie in compagnies:
        polices_qs = Police.objects.filter(
            historiques__id__in=PoliceAssureur.objects.filter(
                compagnie_id=compagnie.id, type_compagnie_id=1
            ).values('historique_police_id')
        ).distinct()

        if not polices_qs.exists():
            continue  # Si aucune police, on passe à la compagnie suivante

        # Ajout du titre de la compagnie
        sheet.append([])
        sheet.append(["", "COMPAGNIE", compagnie.nom])
        sheet.append(headers)

        prime_ht = 0
        prime_ttc = 0
        prime_ht_n = 0
        total_ht = 0
        total_ht_n = 0
        total_ttc = 0

        for police in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=police.id).order_by('-date_du_jour').first()
            # Vérifier si un historique existe
            if dernier_historique:
                annee_actuelle = dernier_historique.date_du_jour.year

                # Trouver la dernière année disponible en excluant l'année actuelle
                derniere_annee_precedente = HistoriquePolice.objects.filter(
                    police_id=police.id,
                    date_du_jour__year__lt=annee_actuelle
                ).aggregate(Max('date_du_jour__year'))['date_du_jour__year__max']

                # Récupérer le dernier historique de cette année trouvée
                if derniere_annee_precedente:
                    historique_annee_precedente = HistoriquePolice.objects.filter(
                        police_id=police.id,
                        date_du_jour__year=derniere_annee_precedente
                    ).order_by('-date_du_jour').first()

                    prime_ht_n = historique_annee_precedente.prime_ht if historique_annee_precedente else 0
                    total_ht_n += prime_ht_n

                else:
                    historique_annee_precedente = None

                # 2. Détermination des primes
                prime_ht = dernier_historique.prime_ht if dernier_historique else 0
                prime_ttc = dernier_historique.prime_ttc if dernier_historique else 0
                total_ht += prime_ht
                total_ttc += prime_ttc

            else:
                historique_annee_precedente = None

            dernier_mouvement = MouvementPolice.objects.filter(police_id=police.id).order_by('-created_at').first()

            date_for_calcul = datetime.today().date()
            n_90_days = date_for_calcul + timedelta(days=90)

            # Détermination du statut
            if dernier_mouvement:
                if dernier_mouvement.date_fin_periode_garantie:
                    if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                        statut = police.etat_police
                    else:
                        difference_jours = (
                                dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (
                                date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                        nombre_total_mois = difference_jours // 30
                        statut = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                else:
                    statut = police.etat_police
            else:
                statut = ''

            sheet.append([
                police.numero,
                police.client.nom if police.client else '',
                police.client.type_personne.libelle if police.client else '',
                police.produit.branche.nom if police.produit.branche else '',
                police.produit.nom if police.produit else '',
                police.date_fin_effet.strftime("%d/%m/%Y") if police.date_fin_effet else '',
                prime_ht_n,
                prime_ht,
                prime_ttc,
                statut
            ])

        # Ajout des totaux pour la compagnie
        sheet.append(["", "", "", "", "", "TOTAL", total_ht_n, total_ht, total_ttc, ""])
        sheet.append([])  # Ligne vide pour séparation

    # Générer le fichier en mémoire
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def add_portefeuille_compagnie(request):
    compagnie_id = request.POST.get('compagnie_id')
    date_requete = request.POST.get('date_requete') or datetime.today().strftime("%d/%m/%Y")

    if compagnie_id == "TOUT":
        compagnies = Compagnie.objects.all()

        if not compagnies.exists():
            return JsonResponse({
                'statut': 0,
                'message': "Aucune compagnie trouvée."
            })

        output = generate_excel_portefeuille_compagnie(compagnies, date_requete)

        # Enregistrement de génération du portefeuille
        analyse_portefeuille = AnalysePortefeuille.objects.create(
            type_portefeuille=TypePortefeuille.ALL_CIE,
            created_at=datetime.now(),
            created_by=request.user
        )

        # Créer un fichier temporaire
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            tmp_file.write(output.getvalue())
            tmp_file_path = tmp_file.name

        # Enregistrer le fichier dans le champ `fichier`
        with open(tmp_file_path, 'rb') as file:
            analyse_portefeuille.fichier.save("Portefeuille_Global.xlsx", File(file))

        # Supprimer le fichier temporaire après l'avoir enregistré
        os.unlink(tmp_file_path)

        return JsonResponse({
            'statut': 1,
            'message': "Portefeuille global généré avec succès !",
            'data': {
                'filename': "Portefeuille_Global.xlsx",
                'file_base64': base64.b64encode(output.getvalue()).decode()
            }
        })

    else:
        compagnie = Compagnie.objects.filter(id=compagnie_id).first()
        polices_qs = Police.objects.filter(
            historiques__id__in=PoliceAssureur.objects.filter(
                compagnie_id=compagnie_id, type_compagnie_id=1
            ).values('historique_police_id')
        ).distinct()

        if not polices_qs.exists():
            return JsonResponse({
                'statut': 0,
                'message': "Aucune police trouvée pour cette compagnie."
            })

        workbook = generate_excel_portefeuille_compagnie([compagnie], date_requete)

        # Enregistrement de génération du portefeuille
        analyse_portefeuille = AnalysePortefeuille.objects.create(
            compagnie=compagnie,
            type_portefeuille=TypePortefeuille.PAR_CIE,
            created_at=datetime.now(),
            created_by=request.user
        )

        # Créer un fichier temporaire
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            tmp_file.write(workbook.getvalue())
            tmp_file_path = tmp_file.name

        # Enregistrer le fichier dans le champ `fichier`
        with open(tmp_file_path, 'rb') as file:
            analyse_portefeuille.fichier.save(f"Portefeuille_{compagnie.nom}.xlsx", File(file))

        # Supprimer le fichier temporaire après l'avoir enregistré
        os.unlink(tmp_file_path)

        return JsonResponse({
            'statut': 1,
            'message': "Portefeuille par compagnie généré avec succès !",
            'data': {
                'filename': f"Portefeuille_{compagnie.nom}.xlsx",
                'file_base64': base64.b64encode(workbook.getvalue()).decode()
            }
        })


# Chargement des polices liées à la compagnie
def get_client_by_compagnie(request):
    compagnie_id = request.GET.get('compagnie_id')
    date_for_calcul = datetime.today().date()
    n_90_days = date_for_calcul + timedelta(days=90)
    total_ht = 0
    total_ttc = 0
    polices_par_compagnie = {}

    if compagnie_id == "TOUT":
        compagnies = Compagnie.objects.all().order_by('nom')

        for compagnie in compagnies:
            polices_qs = Police.objects.filter(
                historiques__id__in=PoliceAssureur.objects.filter(
                    compagnie_id=compagnie.id, type_compagnie_id=1
                ).values('historique_police_id')
            ).distinct()

            polices = []
            for plc in polices_qs:
                dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
                dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

                if dernier_historique:
                    total_ht += dernier_historique.prime_ht
                    total_ttc += dernier_historique.prime_ttc

                # Détermination du statut
                if dernier_mouvement:
                    if dernier_mouvement.date_fin_periode_garantie:
                        if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                            etat_police = plc.etat_police
                        else:
                            difference_jours = (
                                        dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (
                                        date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                            nombre_total_mois = difference_jours // 30
                            etat_police = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                    else:
                        etat_police = plc.etat_police
                else:
                    etat_police = ''

                polices.append({
                    'id': plc.id,
                    'nom': plc.client.nom if plc.client else '',
                    'prenoms': plc.client.prenoms if plc.client else '',
                    'numero': plc.numero,
                    'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else '',
                    'date_creation': plc.created_at.strftime("%d/%m/%Y") if plc.created_at else '',
                    'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                    'statut': etat_police,
                    'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                    'prime_ttc': money_field(dernier_historique.prime_ttc) if dernier_historique else '',
                })

            if polices:
                polices_par_compagnie[compagnie.nom] = polices

    else:
        compagnie = Compagnie.objects.filter(id=compagnie_id).first()
        polices_qs = Police.objects.filter(
            historiques__id__in=PoliceAssureur.objects.filter(
                compagnie_id=compagnie_id, type_compagnie_id=1
            ).values('historique_police_id')
        ).distinct()

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            if dernier_historique:
                total_ht += dernier_historique.prime_ht
                total_ttc += dernier_historique.prime_ttc

            # Détermination du statut
            if dernier_mouvement:
                if dernier_mouvement.date_fin_periode_garantie:
                    if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                        etat_police = plc.etat_police
                    else:
                        difference_jours = (dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                        nombre_total_mois = difference_jours // 30
                        etat_police = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                else:
                    etat_police = plc.etat_police
            else:
                etat_police = ''

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': plc.numero,
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else '',
                'date_creation': plc.created_at.strftime("%d/%m/%Y") if plc.created_at else '',
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'prime_ttc': money_field(dernier_historique.prime_ttc) if dernier_historique else '',
            })

        polices_par_compagnie[compagnie.nom] = polices

    return JsonResponse({'polices_par_compagnie': polices_par_compagnie, 'total_ht': money_field(total_ht), 'total_ttc': money_field(total_ttc)})


# Portefeuille par commercial
def generate_excel_portefeuille_commercial(commercials, date_requete):
    """Génère un fichier Excel unique regroupant les portefeuilles de toutes les commerciaux."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Portefeuille"

    # En-tête du fichier
    sheet.append(["", "DATE DE LA REQUÊTE", date_requete])
    sheet.append([])  # Ligne vide

    headers = [
        "POLICE", "CLIENT", "TYPE DE CLIENT", "BRANCHE", "PRODUIT", "ÉCHÉANCE",
        "PRIME HT EX N-1", "PRIME HT EX N", "PRIME TTC EX N", "STATUT", "COM ENCAISSEE", "COM ATTENDUE"
    ]

    for commercial in commercials:
        polices_qs = Police.objects.filter(
            id__in=Police.objects.filter(
                client__isnull=False,
                historiques__isnull=False,  # Correction ici
                commercial_id=commercial.id
            ).values_list('id', flat=True)
        ).distinct()

        if not polices_qs.exists():
            continue  # Si aucune police, on passe au commercial suivant

        # Ajout du titre du commercial
        sheet.append([])
        sheet.append(["", "COMMERCIAL", commercial.first_name+ ' '+commercial.last_name])
        sheet.append(headers)

        prime_ht = 0
        prime_ht_n = 0
        prime_ttc = 0
        total_ht = 0
        total_ht_n = 0
        total_ttc = 0
        police_com_enc = 0
        police_com_att = 0
        commission_enc = 0
        commission_att = 0

        for police in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=police.id).order_by('-date_du_jour').first()
            # Vérifier si un historique existe
            if dernier_historique:
                annee_actuelle = dernier_historique.date_du_jour.year

                # Trouver la dernière année disponible en excluant l'année actuelle
                derniere_annee_precedente = HistoriquePolice.objects.filter(
                    police_id=police.id,
                    date_du_jour__year__lt=annee_actuelle
                ).aggregate(Max('date_du_jour__year'))['date_du_jour__year__max']

                # Récupérer le dernier historique de cette année trouvée
                if derniere_annee_precedente:
                    historique_annee_precedente = HistoriquePolice.objects.filter(
                        police_id=police.id,
                        date_du_jour__year=derniere_annee_precedente
                    ).order_by('-date_du_jour').first()

                    prime_ht_n = historique_annee_precedente.prime_ht if historique_annee_precedente else 0
                    total_ht_n += prime_ht_n

                else:
                    historique_annee_precedente = None

                # 1. Détermination des commissions attendues
                police_com_att = dernier_historique.commission_courtage if dernier_historique.commission_courtage else 0
                commission_att += police_com_att

                # 2. Détermination des primes
                prime_ht = dernier_historique.prime_ht if dernier_historique else 0
                prime_ttc = dernier_historique.prime_ttc if dernier_historique else 0
                total_ht += prime_ht
                total_ttc += prime_ttc

            else:
                historique_annee_precedente = None

            # 3. Détermination du statut
            dernier_mouvement = MouvementPolice.objects.filter(police_id=police.id).order_by('-created_at').first()
            date_for_calcul = datetime.today().date()
            n_90_days = date_for_calcul + timedelta(days=90)

            if dernier_mouvement:
                if dernier_mouvement.date_fin_periode_garantie:
                    if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                        statut = police.etat_police
                    else:
                        difference_jours = (
                                dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (
                                date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                        nombre_total_mois = difference_jours // 30
                        statut = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                else:
                    statut = police.etat_police
            else:
                statut = ''

            # 4. Détermination des commissions encaissées
            quittances = Quittance.objects.filter(police_id=police.id)
            sum_quittance = 0
            for quittance in quittances:
                sum_reglement = 0
                reglements = Reglement.objects.filter(quittance_id=quittance.id, statut_commission="ENCAISSEE")
                for reglement in reglements:
                    sum_reglement += reglement.montant_com_courtage
                sum_quittance += sum_reglement

            police_com_enc = sum_quittance
            commission_enc += police_com_enc

            sheet.append([
                police.numero,
                police.client.nom if police.client else '',
                police.client.type_personne.libelle if police.client else '',
                police.produit.branche.nom if police.produit.branche else '',
                police.produit.nom if police.produit else '',
                police.date_fin_effet.strftime("%d/%m/%Y") if police.date_fin_effet else '',
                prime_ht_n,
                prime_ht,
                prime_ttc,
                statut,
                police_com_enc,
                police_com_att,
            ])

        # Ajout des totaux pour le commercial
        sheet.append(["", "", "", "", "", "TOTAL", total_ht_n, total_ht, total_ttc, "", commission_enc, commission_att])
        sheet.append([])  # Ligne vide pour séparation

    # Générer le fichier en mémoire
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def add_portefeuille_commercial(request):
    commercial_id = request.POST.get('commercial_id')
    date_requete = request.POST.get('date_requete') or datetime.today().strftime("%d/%m/%Y")

    if commercial_id == "TOUT":
        commercials = User.objects.all()

        if not commercials.exists():
            return JsonResponse({
                'statut': 0,
                'message': "Aucune commercial trouvée."
            })

        output = generate_excel_portefeuille_commercial(commercials, date_requete)

        # Enregistrement de génération du portefeuille
        analyse_portefeuille = AnalysePortefeuille.objects.create(
            type_portefeuille=TypePortefeuille.ALL_COM,
            created_at=datetime.now(),
            created_by=request.user
        )

        # Créer un fichier temporaire
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            tmp_file.write(output.getvalue())
            tmp_file_path = tmp_file.name

        # Enregistrer le fichier dans le champ `fichier`
        with open(tmp_file_path, 'rb') as file:
            analyse_portefeuille.fichier.save("Portefeuille_Global.xlsx", File(file))

        # Supprimer le fichier temporaire après l'avoir enregistré
        os.unlink(tmp_file_path)

        return JsonResponse({
            'statut': 1,
            'message': "Portefeuille global généré avec succès !",
            'data': {
                'filename': "Portefeuille_Global.xlsx",
                'file_base64': base64.b64encode(output.getvalue()).decode()
            }
        })

    else:
        commercial = User.objects.filter(id=commercial_id).first()
        polices_qs = Police.objects.filter(
            id__in=Police.objects.filter(
                client__isnull=False,
                historiques__isnull=False,  # Correction ici
                commercial_id=commercial.id
            ).values_list('id', flat=True)
        ).distinct()

        if not polices_qs.exists():
            return JsonResponse({
                'statut': 0,
                'message': "Aucune police trouvée pour ce commercial."
            })

        workbook = generate_excel_portefeuille_commercial([commercial], date_requete)

        # Enregistrement de génération du portefeuille
        analyse_portefeuille = AnalysePortefeuille.objects.create(
            commercial=commercial,
            type_portefeuille=TypePortefeuille.PAR_COM,
            created_at=datetime.now(),
            created_by=request.user
        )

        # Créer un fichier temporaire
        with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
            tmp_file.write(workbook.getvalue())
            tmp_file_path = tmp_file.name

        # Enregistrer le fichier dans le champ `fichier`
        with open(tmp_file_path, 'rb') as file:
            analyse_portefeuille.fichier.save(f"Portefeuille_{commercial.first_name+'_'+commercial.last_name}.xlsx", File(file))

        # Supprimer le fichier temporaire après l'avoir enregistré
        os.unlink(tmp_file_path)

        return JsonResponse({
            'statut': 1,
            'message': "Portefeuille par commercial généré avec succès !",
            'data': {
                'filename': f"Portefeuille_{commercial.first_name+'_'+commercial.last_name}.xlsx",
                'file_base64': base64.b64encode(workbook.getvalue()).decode()
            }
        })


# Chargement des polices liées au commercial
def get_client_by_commercial(request):
    commercial_id = request.GET.get('commercial_id')
    date_for_calcul = datetime.today().date()
    n_90_days = date_for_calcul + timedelta(days=90)
    total_ht = 0
    total_ttc = 0
    polices_par_commercial = {}

    if commercial_id == "TOUT":
        commercials = User.objects.all().order_by('first_name')

        for commercial in commercials:
            polices_qs = Police.objects.filter(
                id__in=Police.objects.filter(
                    client__isnull=False,
                    historiques__isnull=False,  # Correction ici
                    commercial_id=commercial.id
                ).values_list('id', flat=True)
            ).distinct()

            polices = []
            for plc in polices_qs:
                dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
                dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

                if dernier_historique:
                    total_ht += dernier_historique.prime_ht
                    total_ttc += dernier_historique.prime_ttc

                # Détermination du statut
                if dernier_mouvement:
                    if dernier_mouvement.date_fin_periode_garantie:
                        if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                            etat_police = plc.etat_police
                        else:
                            difference_jours = (
                                        dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (
                                        date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                            nombre_total_mois = difference_jours // 30
                            etat_police = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                    else:
                        etat_police = plc.etat_police
                else:
                    etat_police = ''

                polices.append({
                    'id': plc.id,
                    'nom': plc.client.nom if plc.client else '',
                    'prenoms': plc.client.prenoms if plc.client else '',
                    'numero': plc.numero,
                    'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else '',
                    'date_creation': plc.created_at.strftime("%d/%m/%Y") if plc.created_at else '',
                    'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                    'statut': etat_police,
                    'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                    'prime_ttc': money_field(dernier_historique.prime_ttc) if dernier_historique else '',
                })

            if polices:
                polices_par_commercial[commercial.first_name +' '+ commercial.last_name] = polices

    else:
        commercial = User.objects.filter(id=commercial_id).first()
        polices_qs = Police.objects.filter(
            id__in=Police.objects.filter(
                client__isnull=False,
                historiques__isnull=False,  # Correction ici
                commercial_id=commercial_id
            ).values_list('id', flat=True)
        ).distinct()

        print('polices', polices_qs)

        polices = []
        for plc in polices_qs:
            dernier_historique = HistoriquePolice.objects.filter(police_id=plc.id).order_by('-date_du_jour').first()
            dernier_mouvement = MouvementPolice.objects.filter(police_id=plc.id).order_by('-created_at').first()

            if dernier_historique:
                total_ht += dernier_historique.prime_ht
                total_ttc += dernier_historique.prime_ttc

            # Détermination du statut
            if dernier_mouvement:
                if dernier_mouvement.date_fin_periode_garantie:
                    if n_90_days < dernier_mouvement.date_fin_periode_garantie:
                        etat_police = plc.etat_police
                    else:
                        difference_jours = (
                                    dernier_mouvement.date_fin_periode_garantie - date_for_calcul).days if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else (
                                    date_for_calcul - dernier_mouvement.date_fin_periode_garantie).days
                        nombre_total_mois = difference_jours // 30
                        etat_police = f"A renouveler dans {nombre_total_mois} mois" if dernier_mouvement.date_fin_periode_garantie > date_for_calcul else f"NON renouvelé depuis {nombre_total_mois} mois"
                else:
                    etat_police = plc.etat_police
            else:
                etat_police = ''

            polices.append({
                'id': plc.id,
                'nom': plc.client.nom if plc.client else '',
                'prenoms': plc.client.prenoms if plc.client else '',
                'numero': plc.numero,
                'date_fin_effet': plc.date_fin_effet.strftime("%d/%m/%Y") if plc.date_fin_effet else '',
                'date_creation': plc.created_at.strftime("%d/%m/%Y") if plc.created_at else '',
                'date_resiliation': dernier_mouvement.date_effet.strftime("%d/%m/%Y") if plc.etat_police == "Résilié" and dernier_mouvement else '',
                'statut': etat_police,
                'prime_ht': money_field(dernier_historique.prime_ht) if dernier_historique else '',
                'prime_ttc': money_field(dernier_historique.prime_ttc) if dernier_historique else '',
            })

        polices_par_commercial[commercial.first_name +' '+ commercial.last_name] = polices
    
    return JsonResponse({'polices_par_commercial': polices_par_commercial, 'total_ht': money_field(total_ht),'total_ttc': money_field(total_ttc)})
